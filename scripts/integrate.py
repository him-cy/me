# -*- coding: utf-8 -*-
"""多批次合并：把第二批(不同前缀)接在第一批编号之后, 统一为一个连续编号集合并写表。
例: Wikimedia(01_) + Pexels(P01_) → 连续 01_0001~01_0109。"""
import os, json, sys, shutil
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")

# ===== 参数（改这里） =====
PHOTO_DIR    = r"D:\photos\01_景观规划"
META_MERGED  = os.path.join(PHOTO_DIR, "_pending.json")        # 合并后统一元数据(输出)
META_BATCH1  = os.path.join(PHOTO_DIR, "_pending_wikimedia.json")  # 第一批(已定稿)
META_BATCH2  = os.path.join(PHOTO_DIR, "_pending_pexels.json")    # 第二批(待并入)
PREFIX1      = "01_"
PREFIX2      = "P01_"
BATCH2_START = 76        # 第一批到 75, 第二批从 76 接
EXCEL_PATH   = r"G:\数据采集项目需求文档\数据采集项目需求文档\采集记录表.xlsx"
CATEGORY     = "01"
# ==========================

b1 = json.load(open(META_BATCH1, encoding="utf-8")) if os.path.exists(META_BATCH1) else []
b2 = json.load(open(META_BATCH2, encoding="utf-8")) if os.path.exists(META_BATCH2) else []
d2 = sorted([f for f in os.listdir(PHOTO_DIR) if f.startswith(PREFIX2) and f.endswith(".jpg")],
           key=lambda x: int(x.split("_")[1].split(".")[0]))

mapping = {fn: "%s%04d.jpg" % (PREFIX1, i) for i, fn in enumerate(d2, start=BATCH2_START)}
tmp_of = {}
for fn, new in mapping.items():
    t = "tmp_it_%s" % new
    os.rename(os.path.join(PHOTO_DIR, fn), os.path.join(PHOTO_DIR, t))
    tmp_of[fn] = t
for fn, new in mapping.items():
    os.rename(os.path.join(PHOTO_DIR, tmp_of[fn]), os.path.join(PHOTO_DIR, new))

b2_by_old = {r["file"]: r for r in b2}
for rec in b2:
    new = mapping.get(rec["file"])
    if new:
        rec["file"] = new; rec["image_id"] = new[:-4]
        rec["local_path"] = os.path.join(PHOTO_DIR, new)

def std(r):
    return {"image_id": r["file"][:-4], "file": r["file"], "function_category": CATEGORY,
            "source_type": r.get("source_type", ""), "license": r.get("license", ""),
            "source_url": r.get("source_url", ""), "author": r.get("author", ""),
            "local_path": r.get("local_path", os.path.join(PHOTO_DIR, r["file"]))}

merged = sorted([std(r) for r in b1] + [std(r) for r in b2],
                key=lambda x: int(x["file"].split("_")[1].split(".")[0]))
json.dump(merged, open(META_MERGED, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
if os.path.exists(META_BATCH2):
    shutil.move(META_BATCH2, META_BATCH2 + ".archived")
print("合并元数据:", len(merged), "条 ->", META_MERGED)

# Excel: 在 CATEGORY 区块末插入第二批(已改名)
wb = openpyxl.load_workbook(EXCEL_PATH); ws = wb.active
last_cat = max((r for r in range(3, ws.max_row+1)
                if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith(CATEGORY + "_")),
               default=2)
shutil.copyfile(EXCEL_PATH, EXCEL_PATH + ".bak_before_integrate")
ws.insert_rows(last_cat + 1, len(b2))
b2_sorted = sorted(b2, key=lambda x: int(mapping.get(x["file"], x["file"]).split("_")[1].split(".")[0]))
for i, rec in enumerate(b2_sorted):
    r = last_cat + 1 + i
    img = mapping.get(rec["file"], rec["file"])
    vals = [img, "", "", CATEGORY, "", "", "", rec.get("source_type", "Pexels"),
            rec.get("license", "Pexels License (free to use)"), rec.get("source_url", ""),
            rec.get("author", ""), "", "", "", ""]
    for c, v in enumerate(vals, start=1):
        ws.cell(row=r, column=c, value=v)
wb.save(EXCEL_PATH)
print("Excel 插入:", len(b2), "行")

from collections import Counter
cat = Counter()
for r in range(3, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v: cat[str(v)[:2]] += 1
disk = [f for f in os.listdir(PHOTO_DIR) if f.startswith(PREFIX1) and f.endswith(".jpg")]
print("磁盘:", len(disk), "| 元数据:", len(merged), "| Excel:", dict(cat), "| 一致:", len(disk)==len(merged))
print("PREFIX2 残留:", [f for f in os.listdir(PHOTO_DIR) if f.startswith(PREFIX2)])
