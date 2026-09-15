# -*- coding: utf-8 -*-
"""把 _pending.json 的接受记录写入采集表(对应 CATEGORY 区块)。幂等: 已存在的 image_id 跳过。"""
import os, json, sys, shutil
from _config import cat_dir, TABLE_PATH, ensure
ensure("openpyxl")
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")

# ===== 参数（只改任务相关；路径自动解析，不用手写） =====
CAT_FOLDER  = "01_景观规划"
PHOTO_DIR  = cat_dir(CAT_FOLDER)
META_FILE  = os.path.join(PHOTO_DIR, "_pending.json")
EXCEL_PATH = TABLE_PATH
CATEGORY   = "01"
SOURCE_TYPE = "Pexels"   # 或 "Wikimedia Commons"
# ==========================

records = json.load(open(META_FILE, encoding="utf-8"))
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

# 找 CATEGORY 区块
block_rows = [r for r in range(3, ws.max_row + 1)
              if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith(CATEGORY + "_")]
existing = set(str(ws.cell(row=r, column=1).value) for r in block_rows)
last_cat = max(block_rows) if block_rows else 2
insert_at = last_cat + 1

new = [r for r in records if r["file"] not in existing]
print("待写入(新):", len(new), "| 已存在跳过:", len(records) - len(new))

if new:
    shutil.copyfile(EXCEL_PATH, EXCEL_PATH + ".bak_before_finalize")
    ws.insert_rows(insert_at, len(new))
    for i, rec in enumerate(new):
        r = insert_at + i
        vals = [rec["file"], "", "", CATEGORY, "", "", "",
                rec.get("source_type", SOURCE_TYPE),
                rec.get("license", ""),
                rec.get("source_url", ""),
                rec.get("author", "") or rec.get("meta", {}).get("photographer", ""),
                "", "", "", ""]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(EXCEL_PATH)
    print("已写入 Excel:", len(new), "行于 row", insert_at, "~", insert_at + len(new) - 1)

# 验证
from collections import Counter
cat = Counter()
for r in range(3, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v: cat[str(v)[:2]] += 1
disk = [f for f in os.listdir(PHOTO_DIR) if f.startswith(CATEGORY + "_") and f.endswith(".jpg")]
print("磁盘:", len(disk), "| 元数据:", len(records), "| Excel 分类:", dict(cat))
print("一致:", len(disk) == len(records) == cat.get(CATEGORY, 0))
